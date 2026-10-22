"""Generate real PDF / DOCX bytes for tests (no binary fixtures checked in)."""

from __future__ import annotations

import io

import docx
from fpdf import FPDF


def make_pdf(pages: list[list[str]]) -> bytes:
    """Each page is a list of lines; lines in ALL CAPS render as headings."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    for lines in pages:
        pdf.add_page()
        for line in lines:
            bold = line.isupper()
            pdf.set_font("Helvetica", "B" if bold else "", 14 if bold else 11)
            pdf.multi_cell(0, 6, line)
            pdf.ln(3)
    return bytes(pdf.output())


def make_docx(blocks: list[tuple[str, int]], table: list[list[str]] | None = None) -> bytes:
    """``blocks`` = (text, heading_level); level 0 is a normal paragraph."""
    document = docx.Document()
    for text, level in blocks:
        if level:
            document.add_heading(text, level=level)
        else:
            document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
