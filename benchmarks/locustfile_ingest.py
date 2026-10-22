"""Ingestion throughput test: how many documents per minute go from upload to `ready`.

Each user uploads unique documents (so idempotent dedup never short-circuits) built from
the evaluation corpus: Markdown, PDF and DOCX in equal parts. Locust records the upload
API latency; when the test stops, we wait for the Kafka backlog to drain and compute
end-to-end throughput from the documents' own timestamps:

    docs/min = ready documents / (last processed_at - first created_at) * 60

    locust -f benchmarks/locustfile_ingest.py --host http://localhost:8000 \\
        --headless -u 8 -r 8 -t 1m --csv benchmarks/results/<run>/ingest
"""

from __future__ import annotations

import io
import itertools
import json
import os
import random
import statistics
import time
import uuid
from datetime import datetime
from pathlib import Path

import docx
import httpx
from fpdf import FPDF
from locust import HttpUser, between, events, task
from locust.env import Environment

from benchmarks.common import ROOT, create_tenant

SECTIONS = [
    block.strip()
    for path in sorted((ROOT / "eval" / "corpus_src").glob("*.md"))
    for block in path.read_text().split("\n## ")[1:]
]
_counter = itertools.count()
_headers: dict[str, str] = {}
SUMMARY_PATH = os.environ.get("INGEST_SUMMARY", "")


def make_document(kind: str) -> tuple[str, bytes]:
    marker = f"{uuid.uuid4().hex} {next(_counter)}"
    sections = random.sample(SECTIONS, k=4)
    title = f"Generated document {marker}"
    if kind == "md":
        body = "\n\n".join(f"## {s}" for s in sections)
        return f"{marker}.md", f"# {title}\n\n{body}\n".encode()
    blocks = [(title, 1)] + [
        part for s in sections for part in ((s.split("\n", 1)[0], 2), (s.split("\n", 1)[-1], 0))
    ]
    if kind == "docx":
        document = docx.Document()
        for text, level in blocks:
            if level:
                document.add_heading(text, level=level)
            else:
                document.add_paragraph(" ".join(text.split()))
        buffer = io.BytesIO()
        document.save(buffer)
        return f"{marker}.docx", buffer.getvalue()
    pdf = FPDF()
    pdf.add_page()
    section = 0
    for text, level in blocks:
        pdf.set_font("Helvetica", "B" if level else "", 13 if level else 11)
        safe = " ".join(text.split()).encode("latin-1", "replace").decode("latin-1")
        if level == 2:
            section += 1
            safe = f"{section}. {safe}"  # numbered headings are what PDF parsing can detect
        pdf.multi_cell(0, 6, safe.upper() if level == 1 else safe)
        pdf.ln(2)  # back to the left margin: fpdf2 leaves the cursor at the cell's right edge
    return f"{marker}.pdf", bytes(pdf.output())


@events.test_start.add_listener
def setup(environment: Environment, **_: object) -> None:
    assert environment.host, "--host is required"
    _headers.update(create_tenant(environment.host, upload_corpus=False))


@events.test_stop.add_listener
def report_throughput(environment: Environment, **_: object) -> None:
    assert environment.host
    with httpx.Client(base_url=environment.host, timeout=60) as client:
        deadline = time.monotonic() + 600
        while time.monotonic() < deadline:
            docs = client.get("/v1/documents?limit=200", headers=_headers).json()
            items = docs["items"]
            offset = 200
            while len(items) < docs["total"]:
                items += client.get(
                    f"/v1/documents?limit=200&offset={offset}", headers=_headers
                ).json()["items"]
                offset += 200
            if all(d["status"] in ("ready", "failed") for d in items):
                break
            time.sleep(1)
    ready = [d for d in items if d["status"] == "ready"]
    if not ready:
        print("no documents processed")
        return
    created = [datetime.fromisoformat(d["created_at"]) for d in items]
    processed = [datetime.fromisoformat(d["processed_at"]) for d in ready]
    span = (max(processed) - min(created)).total_seconds()
    latencies = sorted(
        (
            datetime.fromisoformat(d["processed_at"]) - datetime.fromisoformat(d["created_at"])
        ).total_seconds()
        for d in ready
    )
    summary = {
        "documents_uploaded": len(items),
        "documents_ready": len(ready),
        "documents_failed": len(items) - len(ready),
        "chunks": sum(d["chunk_count"] or 0 for d in ready),
        "wall_seconds": round(span, 2),
        "docs_per_minute": round(len(ready) / span * 60, 1),
        "upload_to_ready_seconds_p50": round(statistics.median(latencies), 3),
        "upload_to_ready_seconds_p95": round(
            latencies[max(0, round(0.95 * len(latencies)) - 1)], 3
        ),
        "by_kind": {
            k: sum(1 for d in ready if d["kind"] == k) for k in ("markdown", "pdf", "docx")
        },
    }
    print("INGEST THROUGHPUT", json.dumps(summary))
    if SUMMARY_PATH:
        Path(SUMMARY_PATH).write_text(json.dumps(summary, indent=2) + "\n")


class UploadUser(HttpUser):
    wait_time = between(0.05, 0.2)
    kinds = itertools.cycle(["md", "pdf", "docx"])

    @task
    def upload(self) -> None:
        name, data = make_document(next(self.kinds))
        self.client.post(
            "/v1/documents", files={"file": (name, data)}, headers=_headers, name="upload"
        )
