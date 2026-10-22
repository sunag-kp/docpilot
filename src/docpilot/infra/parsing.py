"""Document parsers: bytes -> ``ParsedDocument`` (headings + paragraphs + page numbers).

Structure matters for chunking and citations: PDFs keep page numbers, DOCX and Markdown
keep heading hierarchy. Any parser failure is *permanent* (retrying a corrupt file is
pointless), so it is raised as ``PermanentIngestionError`` and dead-lettered.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator

import docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from docpilot.domain.errors import PermanentIngestionError
from docpilot.domain.models import Block, DocumentKind, ParsedDocument

_MAX_PDF_PAGES = 2000


class DefaultDocumentParser:
    def parse(self, data: bytes, kind: DocumentKind) -> ParsedDocument:
        try:
            match kind:
                case DocumentKind.PDF:
                    parsed = parse_pdf(data)
                case DocumentKind.DOCX:
                    parsed = parse_docx(data)
                case DocumentKind.MARKDOWN:
                    parsed = parse_markdown(data.decode("utf-8"))
                case DocumentKind.TEXT:
                    parsed = parse_text(data.decode("utf-8"))
        except PermanentIngestionError:
            raise
        except Exception as exc:  # corrupt input surfaces as many library-specific errors
            raise PermanentIngestionError(f"could not parse {kind.value}: {exc}") from exc
        if not any(block.text.strip() and not block.is_heading for block in parsed.blocks):
            raise PermanentIngestionError(
                "document contains no extractable text (scanned PDFs need OCR)"
            )
        return parsed


# --- PDF -------------------------------------------------------------------------------

_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+){0,3})\.?\s+[A-Z][^.!?]{1,80}$")
_SENTENCE_END = re.compile(r"[.!?:;)\"']$")


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise PermanentIngestionError("encrypted PDFs are not supported")
        pages = reader.pages
        if len(pages) > _MAX_PDF_PAGES:
            raise PermanentIngestionError(f"PDF has more than {_MAX_PDF_PAGES} pages")
        blocks: list[Block] = []
        for number, page in enumerate(pages, start=1):
            blocks.extend(_pdf_page_blocks(page.extract_text() or "", number))
    except PyPdfError as exc:
        raise PermanentIngestionError(f"corrupt PDF: {exc}") from exc
    return ParsedDocument(blocks=blocks, page_count=len(pages))


def _pdf_heading_level(line: str) -> int:
    """PDF text has no semantic headings, so detect them conservatively from shape."""
    words = line.split()
    if not 1 <= len(words) <= 10 or _SENTENCE_END.search(line):
        return 0
    numbered = _NUMBERED_HEADING.match(line)
    if numbered:
        return min(numbered.group(1).count(".") + 1, 6)
    letters = [c for c in line if c.isalpha()]
    if len(letters) >= 4 and all(c.isupper() for c in letters):
        return 1
    return 0


def _pdf_page_blocks(text: str, page: int) -> Iterator[Block]:
    paragraph: list[str] = []

    def flush() -> Iterator[Block]:
        if paragraph:
            joined = " ".join(paragraph)
            joined = re.sub(r"(\w)- (\w)", r"\1\2", joined)  # re-join hyphenated line breaks
            yield Block(text=joined, page=page)
            paragraph.clear()

    for raw in text.splitlines():
        line = " ".join(raw.split())
        if not line:
            yield from flush()
            continue
        level = _pdf_heading_level(line)
        if level:
            yield from flush()
            yield Block(text=line, page=page, heading_level=level)
            continue
        paragraph.append(line)
    yield from flush()


# --- DOCX ------------------------------------------------------------------------------

_HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)


def parse_docx(data: bytes) -> ParsedDocument:
    document = docx.Document(io.BytesIO(data))
    blocks: list[Block] = []
    for item in document.iter_inner_content():
        if isinstance(item, DocxParagraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name if item.style is not None else "") or ""
            match = _HEADING_STYLE.match(style)
            level = int(match.group(1)) if match else (1 if style.lower() == "title" else 0)
            blocks.append(Block(text=text, heading_level=level))
        elif isinstance(item, DocxTable):
            for row in item.rows:
                cells = [cell.text.strip() for cell in row.cells]
                # merged cells repeat their text; keep each value once per row
                unique = list(dict.fromkeys(c for c in cells if c))
                if unique:
                    blocks.append(Block(text=" | ".join(unique)))
    return ParsedDocument(blocks=blocks)


# --- Markdown / plain text -------------------------------------------------------------

_ATX_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_SETEXT = re.compile(r"^(=+|-+)\s*$")
_FENCE = re.compile(r"^(```|~~~)")
_IMAGE = re.compile(r"!\[([^\]]*)\]\([^)]*\)")
_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_EMPHASIS = re.compile(r"(\*\*|__|\*|_|~~)(?=\S)(.+?)(?<=\S)\1")
_INLINE_CODE = re.compile(r"`([^`]+)`")
_LIST_MARKER = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
_TABLE_RULE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")


def _strip_inline_markdown(text: str) -> str:
    text = _IMAGE.sub(r"\1", text)
    text = _LINK.sub(r"\1", text)
    text = _INLINE_CODE.sub(r"\1", text)
    text = _EMPHASIS.sub(r"\2", text)
    return text.strip()


def parse_markdown(text: str) -> ParsedDocument:
    lines = text.replace("\r\n", "\n").split("\n")
    if lines and lines[0].strip() == "---":  # YAML front matter
        end = next((i for i, line in enumerate(lines[1:], 1) if line.strip() == "---"), None)
        if end is not None:
            lines = lines[end + 1 :]

    blocks: list[Block] = []
    paragraph: list[str] = []
    in_fence = False

    def flush() -> None:
        if paragraph:
            joined = " ".join(p for p in paragraph if p)
            if joined:
                blocks.append(Block(text=joined))
            paragraph.clear()

    for index, raw in enumerate(lines):
        line = raw.rstrip()
        if _FENCE.match(line.strip()):
            in_fence = not in_fence
            flush()
            continue
        if in_fence:
            paragraph.append(line.strip())
            continue
        heading = _ATX_HEADING.match(line)
        if heading:
            flush()
            blocks.append(
                Block(
                    text=_strip_inline_markdown(heading.group(2)),
                    heading_level=len(heading.group(1)),
                )
            )
            continue
        next_line = lines[index + 1] if index + 1 < len(lines) else ""
        if line.strip() and _SETEXT.match(next_line) and not _LIST_MARKER.match(line):
            flush()
            level = 1 if next_line.strip().startswith("=") else 2
            blocks.append(Block(text=_strip_inline_markdown(line), heading_level=level))
            continue
        if _SETEXT.match(line) or _TABLE_RULE.match(line.strip()):
            continue
        if not line.strip():
            flush()
            continue
        if _LIST_MARKER.match(line):
            # Each list item ends with an implicit full stop so sentences don't run together.
            item = _strip_inline_markdown(_LIST_MARKER.sub("", line))
            paragraph.append(item if item.endswith((".", "!", "?", ":")) else f"{item}.")
            continue
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):  # table row
            cells = [_strip_inline_markdown(c) for c in stripped.strip("|").split("|")]
            paragraph.append(" | ".join(c for c in cells if c) + ".")
            continue
        paragraph.append(_strip_inline_markdown(stripped))
    flush()
    return ParsedDocument(blocks=blocks)


def parse_text(text: str) -> ParsedDocument:
    paragraphs = re.split(r"\n\s*\n", text.replace("\r\n", "\n"))
    return ParsedDocument(blocks=[Block(text=" ".join(p.split())) for p in paragraphs if p.strip()])
