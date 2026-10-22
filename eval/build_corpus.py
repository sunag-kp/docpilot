"""Render the evaluation corpus from Markdown sources into PDF / DOCX / Markdown.

The corpus is fictional (written for this repo), so it is free of licensing issues and
cannot have leaked into any model's training data. Mixing formats makes the evaluation
exercise every parser, not just Markdown.

    uv run python eval/build_corpus.py
"""

from __future__ import annotations

import re
import shutil
from datetime import UTC, datetime
from pathlib import Path

import docx
from fpdf import FPDF

ROOT = Path(__file__).parent
SOURCES = ROOT / "corpus_src"
OUTPUT = ROOT / "corpus"

FORMATS = {
    "api-reference": "pdf",
    "service-level-agreement": "pdf",
    "security-policy": "docx",
    "product-faq": "docx",
}


def parse(markdown: str) -> list[tuple[int, str]]:
    """Return (heading_level, text) blocks; level 0 is a paragraph."""
    blocks: list[tuple[int, str]] = []
    for raw in re.split(r"\n\s*\n", markdown.strip()):
        text = " ".join(raw.split())
        heading = re.match(r"^(#+)\s+(.*)$", text)
        if heading:
            blocks.append((len(heading.group(1)), heading.group(2)))
        else:
            blocks.append((0, text))
    return blocks


def to_pdf(blocks: list[tuple[int, str]], target: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_creation_date(datetime(2026, 1, 1, tzinfo=UTC))  # fixed: reproducible bytes
    pdf.add_page()
    section = 0
    for level, text in blocks:
        if level == 1:
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 9, text.upper())  # ALL CAPS: detected as a heading
        elif level >= 2:
            section += 1
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 13)
            pdf.multi_cell(0, 7, f"{section}. {text}")  # numbered: detected as a heading
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 5.5, text)
        pdf.ln(3)
    pdf.output(str(target))


def to_docx(blocks: list[tuple[int, str]], target: Path) -> None:
    document = docx.Document()
    for level, text in blocks:
        if level:
            document.add_heading(text, level=min(level, 3))
        else:
            document.add_paragraph(text)
    document.core_properties.author = "Northwind Labs"
    document.save(str(target))


def main() -> None:
    OUTPUT.mkdir(exist_ok=True)
    for old in OUTPUT.iterdir():
        old.unlink()
    for source in sorted(SOURCES.glob("*.md")):
        fmt = FORMATS.get(source.stem, "md")
        target = OUTPUT / f"{source.stem}.{fmt}"
        if fmt == "md":
            shutil.copyfile(source, target)
        elif fmt == "pdf":
            to_pdf(parse(source.read_text()), target)
        else:
            to_docx(parse(source.read_text()), target)
        print(f"{target.relative_to(ROOT)}  ({target.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
